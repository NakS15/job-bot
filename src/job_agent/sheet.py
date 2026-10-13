from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import db
from .config import SHEET_PATH, ensure_dirs

COLUMNS = [
    "Approve", "Job ID", "Title", "Company", "Location", "Apply type", "URL",
    "Match score (LLM)", "Original keyword %", "Tailored keyword %",
    "Still missing keywords", "Changes made", "Tailored resume", "Status", "Notes",
]


def read_approvals() -> dict[str, str]:
    """Approve column values keyed by job id, as edited by the user."""
    if not SHEET_PATH.exists():
        return {}
    ws = load_workbook(SHEET_PATH).active
    header = [c.value for c in ws[1]]
    ai, ji = header.index("Approve"), header.index("Job ID")
    return {
        str(r[ji]): str(r[ai] or "").strip().upper()
        for r in ws.iter_rows(min_row=2, values_only=True) if r[ji]
    }


def build_sheet() -> str:
    ensure_dirs()
    approvals = read_approvals()
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")

    for j in db.list_jobs():
        ws.append([
            approvals.get(j["job_id"], ""), j["job_id"], j["title"], j["company"],
            j["location"], j["apply_type"], j["url"], j["match_score"],
            j["original_coverage"], j["tailored_coverage"], ", ".join(j["missing"]),
            j["changes"], j["tailored_pdf"], j["status"], j["notes"],
        ])
        row = ws.max_row
        if j["url"]:
            ws.cell(row, COLUMNS.index("URL") + 1).hyperlink = j["url"]

    dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{max(ws.max_row, 2)}")
    ws.freeze_panes = "C2"
    widths = {"Title": 32, "Company": 24, "URL": 40, "Changes made": 60,
              "Still missing keywords": 36, "Tailored resume": 40, "Notes": 40}
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)
    wb.save(SHEET_PATH)
    return str(SHEET_PATH)
